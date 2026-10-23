#!/usr/bin/env python3
"""Genera el manual Word desde INSTALACION-NATIVA.md; requiere python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE

root = Path(__file__).resolve().parents[1]
source = root / 'docs/INSTALACION-NATIVA.md'
output = root / 'docs/Instalacion_nativa_API_prendas_APA7.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.5)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 2
normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Inches(0.5)
normal.paragraph_format.widow_control = True
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 2
    style.paragraph_format.first_line_indent = Inches(0)
    style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.keep_with_next = True
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == 'Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
code_style = doc.styles.add_style('Codigo', 1)
code_style.base_style = normal
code_style.font.name = 'Courier New'
code_style.font.size = Pt(10)
code_style.paragraph_format.first_line_indent = Inches(0)
code_style.paragraph_format.line_spacing = 2
code_style.paragraph_format.keep_together = True
ref_style = doc.styles.add_style('Referencia APA', 1)
ref_style.base_style = normal
ref_style.paragraph_format.left_indent = Inches(0.5)
ref_style.paragraph_format.first_line_indent = Inches(-0.5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.first_line_indent = Inches(0)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
header._p.append(field)

settings = doc.settings.element
update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'es-ES')
normal.element.get_or_add_rPr().append(lang)

for style in [normal, code_style, ref_style] + [doc.styles[n] for n in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']]:
    fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
    if fonts is not None:
        for key in list(fonts.attrib):
            if 'Theme' in key:
                del fonts.attrib[key]
        for key in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            fonts.set(qn('w:' + key), 'Courier New' if style.name == 'Codigo' else 'Times New Roman')

# Portada técnica: no se inventan autor personal, institución, asignatura ni docente.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(72)
p.paragraph_format.first_line_indent = Inches(0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Instalación nativa del modelo entrenado y su API de reconocimiento de prendas'); r.bold = True
p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Inches(0)
for text in ['Manual técnico de instalación', 'Proyecto modelo-reconocimiento-prendas', 'API 2.0.0 · MobileNetV3-Small · CPU', 'Edición revisada 1.1', '23 de octubre de 2026']:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()
doc.add_paragraph('Instalación nativa del modelo entrenado y su API de reconocimiento de prendas', 'Heading 1')


def hyperlink(paragraph, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
    run.append(props); txt = OxmlElement('w:t'); txt.text = url; run.append(txt)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, text):
    for token in re.split(r'(`[^`]+`|\*[^*]+\*|https?://\S+)', text):
        if not token:
            continue
        if token.startswith('http'):
            hyperlink(paragraph, token)
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1]); r.font.name = 'Courier New'; r.font.size = Pt(10)
        elif token.startswith('*'):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)

lines = source.read_text().splitlines()
body = '\n'.join(lines[lines.index('## Objetivo y alcance'):])
in_code = False
references = False
pending = []
table_lines = []

def flush():
    if pending:
        p = doc.add_paragraph(style='Referencia APA' if references else 'Normal')
        inline(p, ' '.join(pending))
        pending.clear()

def render_table():
    if not table_lines:
        return
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')]
            for line in table_lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.autofit = False
    for col, width in zip(table.columns, [1.4, 2.0, 3.1]):
        col.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'single' if edge in ['top', 'bottom'] else 'nil')
        border.set(qn('w:sz'), '6')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    props.append(borders)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        row_props = table.rows[-1]._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); row_props.append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); row_props.append(repeat)
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = False
            inline(p, value)
            if index == 0:
                for run in p.runs:
                    run.bold = True
                cell_borders = OxmlElement('w:tcBorders')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
                cell_borders.append(bottom); cell._tc.get_or_add_tcPr().append(cell_borders)
    table_lines.clear()


for line in body.splitlines():
    if not in_code and line.startswith('|'):
        flush()
        table_lines.append(line)
        continue
    if table_lines:
        render_table()
    if line.startswith('```'):
        flush()
        if in_code:
            doc.paragraphs[-1].paragraph_format.keep_with_next = False
        in_code = not in_code
        continue
    if in_code:
        p = doc.add_paragraph(line if line else ' ', 'Codigo')
        p.paragraph_format.keep_with_next = True
    elif line == 'Tabla 1':
        flush()
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        p.add_run(line).bold = True
    elif line.startswith('*Diagnóstico de incidencias'):
        flush()
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.keep_with_next = True
        inline(p, line)
    elif line.startswith('## '):
        flush()
        title = line[3:]
        if title == 'Referencias':
            references = True
            doc.add_page_break()
        doc.add_paragraph(title, 'Heading 1')
    elif not line.strip():
        flush()
    else:
        pending.append(line)
flush()
render_table()
doc.core_properties.title = 'Instalación nativa del modelo entrenado y su API de reconocimiento de prendas'
doc.core_properties.subject = 'Despliegue directo con Python; API de una imagen; formato APA 7'
doc.core_properties.author = ''
doc.core_properties.keywords = 'APA 7, instalación, API, MobileNetV3-Small, CPU, una imagen'
doc.core_properties.language = 'es-ES'
doc.save(output)
print(output)
print('Parrafos:', len(doc.paragraphs))
